"""Formatted Excel workbooks with automatic best-value highlighting.

CSV files remain the raw source of truth; these workbooks are convenience views.
Uses openpyxl only (no Excel / LibreOffice dependency). Per-sheet, per-workbook
failures are caught so one bad sheet never aborts aggregation.

Highlighting policy (method-comparison tables): within each comparison group
(all non-metric dimension columns except ``method_name``) the best metric value
is green, second light-green, third pale-yellow. Direction-aware via the
HIGHER_BETTER / LOWER_BETTER registries.
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

try:  # pragma: no cover - presence depends on the environment
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.dataframe import dataframe_to_rows
    _OPENPYXL = True
except Exception:  # pragma: no cover
    openpyxl = None
    _OPENPYXL = False

HIGHER_BETTER = {
    "mean_ari", "median_ari", "max_ari", "mean_nmi", "mean_ami", "mean_v_measure",
    "mean_homogeneity", "mean_completeness", "mean_fowlkes_mallows", "mean_purity",
    "k_accuracy", "success_rate", "selected_is_trace_best_rate", "win_rate",
    "mean_trace_best_ari", "median_trace_best_ari", "exact_k_rate",
}
LOWER_BETTER = {
    "failure_rate", "mean_runtime_sec", "median_runtime_sec", "mean_k_abs_error",
    "median_k_abs_error", "mean_ari_gap_to_trace_best", "median_ari_gap_to_trace_best",
    "mean_relative_ari_gap", "median_relative_ari_gap", "loss_rate",
    "mean_selected_rank_by_ari", "mean_trace_best_rank_by_internal_objective",
}

_FILL_BEST = "63BE7B"
_FILL_2ND = "C6EFCE"
_FILL_3RD = "FFEB9C"

_PCT_COLS = {
    "success_rate", "failure_rate", "k_accuracy", "win_rate", "tie_rate", "loss_rate",
    "selected_is_trace_best_rate", "overcluster_rate", "undercluster_rate",
    "exact_k_rate",
}
_INT_COLS = {"count", "n", "n_paired", "n_success", "n_with_trace", "top_k",
             "true_k", "best_selected_k", "selected_k", "cluster_count"}


def available() -> bool:
    return _OPENPYXL


def safe_sheet_name(name: str, used: set) -> str:
    """Excel-safe (<=31 char, no illegal chars), de-duplicated sheet name."""
    safe = "".join(ch for ch in str(name) if ch not in '[]:*?/\\')[:31] or "sheet"
    base = safe
    i = 1
    while safe in used:
        suffix = f"~{i}"
        safe = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(safe)
    return safe


def _number_format(col: str) -> Optional[str]:
    if col in _INT_COLS or col.startswith("rank_"):
        return "0"
    if col in _PCT_COLS or col.endswith("_rate"):
        return "0.0%"
    if "runtime" in col:
        return "0.00"
    lowered = col.lower()
    if any(tok in lowered for tok in (
        "ari", "nmi", "ami", "measure", "homogeneity", "completeness",
        "fowlkes", "purity", "gap", "weight", "entropy", "score", "diff", "p_",
        "cohens", "_p", "trace_best", "rank",
    )):
        return "0.0000"
    return None


def _direction(col: str) -> Optional[str]:
    if col in HIGHER_BETTER:
        return "higher"
    if col in LOWER_BETTER:
        return "lower"
    return None


RANK_SPECS = {
    "rank_mean_ari": ("mean_ari", "higher"),
    "rank_k_accuracy": ("k_accuracy", "higher"),
    "rank_runtime": ("mean_runtime_sec", "lower"),
    "rank_trace_best_gap": ("mean_ari_gap_to_trace_best", "lower"),
}


def add_rank_columns(df: pd.DataFrame, comparison_group_cols: Sequence[str]) -> pd.DataFrame:
    """Add within-group rank columns for key metrics (rank 1 = best)."""
    if df.empty:
        return df
    out = df.copy()
    gcols = [c for c in comparison_group_cols if c in out.columns]
    for rank_col, (src, direction) in RANK_SPECS.items():
        if src not in out.columns:
            continue
        vals = pd.to_numeric(out[src], errors="coerce")
        ascending = (direction == "lower")
        if gcols:
            out[rank_col] = vals.groupby([out[c] for c in gcols]).rank(
                ascending=ascending, method="min")
        else:
            out[rank_col] = vals.rank(ascending=ascending, method="min")
    return out


def _method_col(df: pd.DataFrame) -> Optional[str]:
    if "method_name" in df.columns:
        return "method_name"
    if "method" in df.columns:
        return "method"
    return None


def _autosize(ws) -> None:
    for col_cells in ws.columns:
        width = 8
        letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            v = cell.value
            if v is not None:
                width = max(width, min(40, len(str(v)) + 2))
        ws.column_dimensions[letter].width = width


def add_sheet(
    wb,
    sheet_name: str,
    df: pd.DataFrame,
    *,
    comparison_group_cols: Optional[Sequence[str]] = None,
    highlight: bool = True,
    add_ranks: bool = True,
    used_names: Optional[set] = None,
) -> None:
    if not _OPENPYXL:
        return
    used_names = used_names if used_names is not None else set()
    name = safe_sheet_name(sheet_name, used_names)
    ws = wb.create_sheet(title=name)
    if df is None or df.empty:
        ws["A1"] = "(empty)"
        return

    gcols = [c for c in (comparison_group_cols or []) if c in df.columns]
    work = df.copy()
    method_col = _method_col(work)
    if add_ranks and method_col == "method_name":
        work = add_rank_columns(work, gcols)

    # Write rows.
    for row in dataframe_to_rows(work, index=False, header=True):
        ws.append(row)

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Number formats.
    col_index = {c: i + 1 for i, c in enumerate(work.columns)}
    for col, idx in col_index.items():
        fmt = _number_format(col)
        if not fmt:
            continue
        letter = get_column_letter(idx)
        for cell in ws[letter][1:]:
            cell.number_format = fmt

    # Highlighting (method-comparison tables only).
    if highlight and method_col is not None:
        _highlight(ws, work, gcols, col_index)

    _autosize(ws)


def _highlight(ws, df: pd.DataFrame, gcols: List[str], col_index: Dict[str, int]) -> None:
    fills = {
        1: PatternFill(start_color=_FILL_BEST, end_color=_FILL_BEST, fill_type="solid"),
        2: PatternFill(start_color=_FILL_2ND, end_color=_FILL_2ND, fill_type="solid"),
        3: PatternFill(start_color=_FILL_3RD, end_color=_FILL_3RD, fill_type="solid"),
    }
    df = df.reset_index(drop=True)
    if gcols:
        groups = df.groupby(gcols, dropna=False).groups
    else:
        groups = {"__all__": df.index}
    for col, idx in col_index.items():
        direction = _direction(col)
        if not direction:
            continue
        letter = get_column_letter(idx)
        for _, members in groups.items():
            members = list(members)
            vals = pd.to_numeric(df.loc[members, col], errors="coerce")
            if vals.notna().sum() == 0:
                continue
            ranks = vals.rank(ascending=(direction == "lower"), method="min")
            for pos in members:
                r = ranks.loc[pos]
                if pd.isna(r):
                    continue
                ri = int(r)
                if ri in fills:
                    ws[f"{letter}{pos + 2}"].fill = fills[ri]


def write_workbook(
    path: Path,
    sheets: List[Tuple[str, pd.DataFrame, Optional[Sequence[str]], bool]],
) -> Optional[str]:
    """Write one workbook. ``sheets`` = [(name, df, comparison_group_cols, highlight)]."""
    if not _OPENPYXL:
        return None
    try:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        used: set = set()
        for name, df, gcols, hl in sheets:
            try:
                add_sheet(wb, name, df, comparison_group_cols=gcols,
                          highlight=hl, used_names=used)
            except Exception as exc:  # one bad sheet never aborts the file
                print(f"[excel] skip sheet {name} in {path.name}: {exc}", flush=True)
        if not wb.sheetnames:
            wb.create_sheet("empty")
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        return str(path)
    except Exception as exc:
        print(f"[excel] skip workbook {path.name}: {exc}", flush=True)
        return None
