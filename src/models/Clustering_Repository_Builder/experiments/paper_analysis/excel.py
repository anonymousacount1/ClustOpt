"""Formatted Excel workbooks (openpyxl only).

CSV remains the source of truth; a workbook is a *readable view* of CSVs that are
already on disk. Every workbook therefore gets, without exception:

* a **Definitions** sheet (terminology + primary metrics + oracle levels),
* a **Method Labels** sheet (short label -> concrete method name),
* a **Selection Manifest** sheet when representatives are involved,
* a **Capability Warnings** sheet when any analysis was skipped,
* frozen header rows, auto-filters, sensible column widths,
* direction-aware conditional formatting (higher-is-better vs lower-is-better),
* explicit number formats for shares (%) and runtimes (s).

No sheet contains hidden selection logic: whatever is displayed comes from the
frames the analysis already wrote.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import paths
from .method_registry import label_map_records

MAX_SHEET_NAME = 31
MAX_ROWS_PER_SHEET = 60000        # keep workbooks openable; overflow noted in-sheet

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
NOTE_FONT = Font(italic=True, color="555555", size=9)
TITLE_FONT = Font(bold=True, size=12)

# Columns where a HIGHER value is better (green at the top of the colour scale).
HIGHER_IS_BETTER: Tuple[str, ...] = (
    "best_view_mean_ari", "best_view_median_ari", "best_view_max_ari",
    "best_view_q25_ari", "best_view_q75_ari", "best_view_k_accuracy",
    "best_view_exact_k_rate", "all_views_mean_ari", "all_views_median_ari",
    "success_rate", "n_success", "win_rate", "mean_ari_gain", "median_ari_gain",
    "selection_rate_per_run", "share_of_selected_slots", "top_metric_rate",
    "dataset_coverage", "weight_mass_share", "selected_slot_share",
    "vbs_best_view_mean_ari", "share_of_full_vbs", "share_of_achievable_gain",
    "selected_is_trace_best_rate", "mean_new_metric_weight_mass",
    "share_runs_with_new_metric", "selection_stability", "unique_coverage_share",
    "headroom_paired_mean", "headroom_unpaired", "marginal_gain",
)
# Columns where a LOWER value is better.
LOWER_IS_BETTER: Tuple[str, ...] = (
    "best_view_std_ari", "best_view_mean_k_abs_error",
    "best_view_median_k_abs_error", "failure_rate", "loss_rate",
    "mean_total_runtime_all_views_sec", "median_total_runtime_all_views_sec",
    "std_total_runtime_all_views_sec", "p90_total_runtime_all_views_sec",
    "sum_total_runtime_all_views_sec", "mean_selected_best_view_runtime_sec",
    "median_selected_best_view_runtime_sec",
    "total_runtime_all_views_mean_sec", "total_runtime_all_views_median_sec",
    "total_runtime_all_views_sum_sec", "best_view_runtime_mean_sec",
    "best_view_runtime_median_sec",
    "portfolio_total_execution_runtime_mean_sec",
    "portfolio_total_execution_runtime_median_sec",
    "portfolio_total_execution_runtime_total_sec",
    "selected_member_total_all_views_runtime_mean_sec",
    "selected_member_total_all_views_runtime_median_sec",
    "search_selection_gap_mean", "search_selection_gap_median",
    "view_gap_mean", "view_gap_median", "wilcoxon_p", "ttest_p", "holm_p",
)
PERCENT_COLUMNS: Tuple[str, ...] = (
    "success_rate", "failure_rate", "win_rate", "tie_rate", "loss_rate",
    "best_view_k_accuracy", "best_view_exact_k_rate",
    "best_view_overcluster_rate", "best_view_undercluster_rate",
    "selection_rate_per_run", "share_of_selected_slots", "top_metric_rate",
    "dataset_coverage", "weight_mass_share", "selected_slot_share",
    "top_metric_share", "share_of_full_vbs", "share_of_achievable_gain",
    "selected_is_trace_best_rate", "share_runs_with_new_metric",
    "share_runs_top_metric_is_new", "mean_new_metric_fraction",
    "median_new_metric_fraction", "share_a_better", "share_b_better",
    "share_tied", "share_reference_better", "share_comparison_better",
    "share_equal", "share_delta_positive", "selection_stability",
    "unique_coverage_share", "win_share", "divergence_rate",
    "share_datasets_vbs_strictly_better", "share_datasets_equal",
)
RUNTIME_COLUMNS_SUFFIX = ("_sec",)
ARI_DECIMALS = 4


@dataclass
class Sheet:
    """One worksheet: a frame plus how to present it."""

    name: str
    frame: Optional[pd.DataFrame]
    description: str = ""
    freeze: str = "A2"
    autofilter: bool = True
    highlight: bool = True
    max_col_width: int = 52


@dataclass
class WorkbookSpec:
    """One workbook: filename, title and its ordered sheets."""

    filename: str
    title: str
    description: str = ""
    sheets: List[Sheet] = field(default_factory=list)

    def add(self, name: str, frame: Optional[pd.DataFrame], description: str = "",
            **kwargs) -> "WorkbookSpec":
        self.sheets.append(Sheet(name, frame, description, **kwargs))
        return self


# --------------------------------------------------------------------------- #
# Standard sheets present in every workbook
# --------------------------------------------------------------------------- #
DEFINITIONS: Tuple[Tuple[str, str], ...] = (
    ("fixed evaluated variant",
     "A concrete method run exactly as defined (e.g. regressor_top5_raw). "
     "Operational as a fixed algorithmic configuration."),
    ("retrospective global representative",
     "The fixed variant with the highest global Mean Best-View ARI on split 1 "
     "inside a predefined candidate pool. It is the BEST OBSERVED single variant "
     "on split 1 -- NOT an independently selected deployment policy."),
    ("retrospective family representative",
     "One fixed variant selected separately inside each structural family. A "
     "family-level retrospective policy ORACLE, reported as headroom only."),
    ("dataset-level VBS",
     "Virtual Best Solver: picks a different candidate per dataset using the "
     "dataset's true external ARI. A non-deployable oracle upper bound."),
    ("Best-View",
     "The highest-ARI of the three views (x_only, y_only, xy_2d). A VIEW ORACLE: "
     "it uses external ARI to pick the representation. All three views must be "
     "run to obtain one, so total_runtime_all_views_sec is the primary cost."),
    ("best_view_mean_ari",
     "PRIMARY METRIC. Mean over datasets of each dataset's Best-View ARI."),
    ("total_runtime_all_views_sec",
     "PRIMARY COST. x_only + y_only + xy_2d runtime for one dataset x method."),
    ("best_view_runtime_sec",
     "DIAGNOSTIC ONLY. Runtime of the single view whose ARI was selected; not the "
     "cost of obtaining a Best-View result."),
    ("selected_member_total_all_views_runtime_sec",
     "VBS DIAGNOSTIC: all-views runtime of the member the oracle picked after "
     "the fact. NOT the cost of obtaining the VBS result."),
    ("portfolio_total_execution_runtime_sec",
     "What realising a VBS actually costs: the sum over EVERY portfolio member "
     "across all required views. The only quantity that may be described as the "
     "cost of a VBS."),
    ("fully_operational",
     "TRUE only when every portfolio member is runnable without test-time "
     "ground truth. The 12-variant MLP/KNN portfolios are FALSE because they "
     "include Dynamic Real-K."),
    ("oracle_contamination_reason",
     "Why a portfolio is not fully operational: real_utility and/or "
     "dynamic_realK."),
    ("selection_stability",
     "Share of bootstrap resamples in which the same candidate was re-selected. "
     "A low value means the retrospective representative is a fragile choice."),
    ("utility oracle",
     "real_* methods use the REAL utility vector. Oracle-level: utility_oracle."),
    ("dynamic-K oracle assistance",
     "*_dynamic_realK_* MLP/KNN methods rank metrics with the learned predictor "
     "but take the metric COUNT from the real utility vector. Oracle-assisted "
     "diagnostics, excluded from the primary meta-learning comparison."),
    ("post_selection_exploratory",
     "A statistical test where at least one arm is a retrospectively selected "
     "representative tested on the same split-1 data. Optimistic by "
     "construction; NOT confirmatory evidence."),
    ("is_independently_preselected",
     "Always FALSE in this analysis: downstream clustering experiments exist for "
     "split 1 only, so no policy could be chosen on independent data."),
)

PAPER_CLAIM_LEVELS: Tuple[Tuple[str, str], ...] = (
    ("fixed_observed", "A fixed evaluated variant's own measured performance."),
    ("retrospective_summary",
     "A retrospectively selected representative: a benchmark summary, not an "
     "independently selected deployment comparison."),
    ("diagnostic_oracle",
     "Uses an oracle (real utility, real metric count, per-family reselection) to "
     "diagnose where performance is lost."),
    ("upper_bound", "A dataset-level VBS: non-deployable upper bound."),
)


def _definitions_frame() -> pd.DataFrame:
    rows = [{"term": t, "definition": d, "kind": "terminology"}
            for t, d in DEFINITIONS]
    rows += [{"term": t, "definition": d, "kind": "paper_claim_level"}
             for t, d in PAPER_CLAIM_LEVELS]
    return pd.DataFrame(rows)


def _label_frame() -> pd.DataFrame:
    return pd.DataFrame(label_map_records())


# --------------------------------------------------------------------------- #
# Writing
# --------------------------------------------------------------------------- #
def _safe_sheet_name(name: str, used: set) -> str:
    clean = "".join(c for c in str(name) if c not in "[]:*?/\\")[:MAX_SHEET_NAME]
    clean = clean or "sheet"
    candidate, i = clean, 2
    while candidate in used:
        suffix = f"_{i}"
        candidate = clean[: MAX_SHEET_NAME - len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate


def _flatten(frame: pd.DataFrame) -> pd.DataFrame:
    out = frame.copy()
    if isinstance(out.index, pd.MultiIndex) or out.index.name is not None:
        out = out.reset_index()
    out.columns = [str(c) for c in out.columns]
    # openpyxl cannot write list/dict cells.
    for col in out.columns:
        if out[col].map(lambda v: isinstance(v, (list, dict, set, tuple))).any():
            out[col] = out[col].map(
                lambda v: ";".join(map(str, v)) if isinstance(v, (list, set, tuple))
                else (str(v) if isinstance(v, dict) else v))
    return out


def _number_format(col: str) -> Optional[str]:
    if col in PERCENT_COLUMNS:
        return "0.0%"
    if col.endswith(RUNTIME_COLUMNS_SUFFIX):
        return "#,##0.0"
    if "_ari" in col or col.endswith(("_gap", "_delta", "_p", "_dz", "_r", "_rho")):
        return "0." + "0" * ARI_DECIMALS
    return None


def _write_sheet(ws: Worksheet, sheet: Sheet) -> None:
    frame = sheet.frame
    row_cursor = 1
    if sheet.description:
        ws.cell(row=1, column=1, value=sheet.description).font = NOTE_FONT
        ws.cell(row=1, column=1).alignment = Alignment(wrap_text=False)
        row_cursor = 3
    if frame is None or frame.empty:
        ws.cell(row=row_cursor, column=1,
                value="(no rows -- this analysis produced no data; see the "
                      "capability warnings sheet)").font = NOTE_FONT
        return

    data = _flatten(frame)
    truncated = False
    if len(data) > MAX_ROWS_PER_SHEET:
        data = data.head(MAX_ROWS_PER_SHEET)
        truncated = True

    header_row = row_cursor
    for j, col in enumerate(data.columns, start=1):
        cell = ws.cell(row=header_row, column=j, value=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    for i, (_, r) in enumerate(data.iterrows(), start=header_row + 1):
        for j, col in enumerate(data.columns, start=1):
            value = r[col]
            if isinstance(value, (np.integer,)):
                value = int(value)
            elif isinstance(value, (np.floating,)):
                value = None if not np.isfinite(value) else float(value)
            elif isinstance(value, (np.bool_,)):
                value = bool(value)
            elif value is pd.NA:
                value = None
            elif isinstance(value, float) and not np.isfinite(value):
                value = None
            ws.cell(row=i, column=j, value=value)

    last_row = header_row + len(data)
    last_col_letter = get_column_letter(len(data.columns))
    if sheet.autofilter:
        ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_row}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    for j, col in enumerate(data.columns, start=1):
        letter = get_column_letter(j)
        fmt = _number_format(col)
        if fmt:
            for i in range(header_row + 1, last_row + 1):
                ws.cell(row=i, column=j).number_format = fmt
        sample = data[col].astype(str).head(200)
        width = max(len(col) + 3, int(sample.str.len().max() or 0) + 2)
        ws.column_dimensions[letter].width = min(sheet.max_col_width, max(9, width))

        if sheet.highlight and len(data) > 1:
            rng = f"{letter}{header_row + 1}:{letter}{last_row}"
            if col in HIGHER_IS_BETTER:
                ws.conditional_formatting.add(rng, ColorScaleRule(
                    start_type="min", start_color="F8696B",
                    mid_type="percentile", mid_value=50, mid_color="FFEB84",
                    end_type="max", end_color="63BE7B"))
            elif col in LOWER_IS_BETTER:
                ws.conditional_formatting.add(rng, ColorScaleRule(
                    start_type="min", start_color="63BE7B",
                    mid_type="percentile", mid_value=50, mid_color="FFEB84",
                    end_type="max", end_color="F8696B"))
    if truncated:
        ws.cell(row=last_row + 2, column=1,
                value=f"NOTE: truncated to the first {MAX_ROWS_PER_SHEET} rows; "
                      f"the full table is in the corresponding CSV.").font = NOTE_FONT


def write_workbook(spec: WorkbookSpec, output_dir: Path, *,
                   selection_manifest: Optional[pd.DataFrame] = None,
                   portfolio_manifest: Optional[pd.DataFrame] = None,
                   capability_warnings: Optional[pd.DataFrame] = None,
                   extra_notes: Sequence[str] = (),
                   dry_run: bool = False) -> Optional[Path]:
    """Write one workbook; returns the path (``None`` when ``dry_run``)."""
    path = Path(output_dir) / spec.filename
    if dry_run:
        return None

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()

    # --- README sheet ---------------------------------------------------- #
    ws = wb.create_sheet(_safe_sheet_name("README", used))
    ws.cell(row=1, column=1, value=spec.title).font = TITLE_FONT
    row = 3
    for line in ([spec.description] if spec.description else []) + list(extra_notes):
        ws.cell(row=row, column=1, value=line).font = NOTE_FONT
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Sheets in this workbook").font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=1, value="sheet").font = HEADER_FONT
    ws.cell(row=row, column=1).fill = HEADER_FILL
    ws.cell(row=row, column=2, value="contents").font = HEADER_FONT
    ws.cell(row=row, column=2).fill = HEADER_FILL
    ws.cell(row=row, column=3, value="rows").font = HEADER_FONT
    ws.cell(row=row, column=3).fill = HEADER_FILL
    row += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 96
    ws.column_dimensions["C"].width = 10

    planned: List[Tuple[str, Sheet]] = []
    for sheet in spec.sheets:
        planned.append((_safe_sheet_name(sheet.name, used), sheet))
    standard = [
        (_safe_sheet_name("Definitions", used),
         Sheet("Definitions", _definitions_frame(),
               "Terminology contract. Read this before quoting any number.",
               highlight=False)),
        (_safe_sheet_name("Method Labels", used),
         Sheet("Method Labels", _label_frame(),
               "Short plot label -> concrete method name and its taxonomy.",
               highlight=False)),
    ]
    if selection_manifest is not None:
        standard.append((_safe_sheet_name("Selection Manifest", used),
                         Sheet("Selection Manifest", selection_manifest,
                               "The FROZEN retrospective selections every table "
                               "in this workbook relies on.", highlight=False)))
    if portfolio_manifest is not None:
        standard.append((_safe_sheet_name("Portfolio Manifest", used),
                         Sheet("Portfolio Manifest", portfolio_manifest,
                               "FROZEN VBS portfolio membership with the "
                               "operational status and oracle-contamination "
                               "reason of each portfolio.", highlight=False)))
    if capability_warnings is not None:
        standard.append((_safe_sheet_name("Capability Warnings", used),
                         Sheet("Capability Warnings", capability_warnings,
                               "Sub-analyses skipped because a method's outputs "
                               "cannot support them, with the exact missing "
                               "field.", highlight=False)))

    for name, sheet in planned + standard:
        n_rows = 0 if sheet.frame is None else len(sheet.frame)
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=sheet.description)
        ws.cell(row=row, column=3, value=n_rows)
        row += 1

    for name, sheet in planned + standard:
        _write_sheet(wb.create_sheet(name), sheet)

    paths.mkdirs(path.parent)
    wb.save(paths.ext(path))
    return path


# --------------------------------------------------------------------------- #
# Preliminary workbook (the only one executed in this stage)
# --------------------------------------------------------------------------- #
def preliminary_workbook(tables: Dict[str, Optional[pd.DataFrame]]) -> WorkbookSpec:
    """``01_preliminary_selection.xlsx`` -- the frozen selection study."""
    spec = WorkbookSpec(
        filename="01_preliminary_selection.xlsx",
        title="Preliminary retrospective representative-selection study (Split 1)",
        description=(
            "Frozen inputs for comparison suites 1-5. Every representative here "
            "was selected RETROSPECTIVELY on split 1 from an explicit candidate "
            "pool; none is an independently preselected deployment policy."),
    )
    spec.add("Global Representatives", tables.get("global_representatives"),
             "One retrospective global representative per candidate pool, with "
             "its selection metric values and any tie-breaks invoked.")
    spec.add("Candidate Rankings", tables.get("candidate_rankings"),
             "Full ranked candidate table behind every global selection "
             "(runners-up included).")
    spec.add("C1 Helper Selections", tables.get("comparison_01_helper"),
             "Comparison 1: best observed fixed Real Top-K raw variant. "
             "real_top_dynamic_raw is recorded separately, never merged in.")
    spec.add("C2 Global Reps", tables.get("comparison_02_global"),
             "Comparison 2: the 15 primary (utility source x metric-count "
             "policy) raw-vs-softmax representatives, plus the oracle-assisted "
             "Dynamic Real-K diagnostics.")
    spec.add("C2 Family Reps", tables.get("comparison_02_family"),
             "Comparison 2: family-level raw-vs-softmax reselection. HEADROOM "
             "ONLY -- a family-level retrospective policy oracle.")
    spec.add("C2 Family Divergence", tables.get("family_divergence"),
             "How often the family-level choice differs from the global one.")
    spec.add("C4 Fixed Reps", tables.get("comparison_04_fixed"),
             "Comparison 4: best observed FIXED MLP and KNN variants, plus the "
             "supplemental pools that also include Dynamic Predict-K.")
    spec.add("Portfolio Manifest", tables.get("portfolio_manifest"),
             "Frozen VBS portfolio membership, with real-utility and "
             "oracle-assisted contamination flags.")
    spec.add("Method Summary", tables.get("method_summary"),
             "Primary-metric block for every one of the 51 collected methods.")
    spec.add("Coverage Audit", tables.get("coverage_audit"),
             "Per-method dataset/view/success coverage on split 1.")
    spec.add("Capability Matrix", tables.get("capability_matrix"),
             "Verified per-method capability matrix.")
    spec.add("Metric Registry", tables.get("metric_registry"),
             "Canonical metric taxonomy (category derived from the "
             "implementation package).")
    spec.add("Selection Aware Bootstrap", tables.get("selection_aware_bootstrap"),
             "Bootstrap that re-selects the representative inside every "
             "resample; wider than a naive interval by design.")
    return spec


# --------------------------------------------------------------------------- #
# Comparison-suite workbooks (implemented, executed only after approval)
# --------------------------------------------------------------------------- #
def comparison_workbook(comparison_id: str, filename: str, title: str,
                        description: str,
                        tables: Dict[str, Optional[pd.DataFrame]],
                        sheet_plan: Sequence[Tuple[str, str, str]]
                        ) -> WorkbookSpec:
    """Generic comparison workbook builder.

    ``sheet_plan`` is a sequence of ``(sheet name, tables key, description)``, so
    each suite declares its own sheets without duplicating the writer.
    """
    spec = WorkbookSpec(filename=filename, title=title, description=description)
    for sheet_name, key, desc in sheet_plan:
        spec.add(sheet_name, tables.get(key), desc)
    return spec


WORKBOOK_FILENAMES: Dict[str, str] = {
    "preliminary": "01_preliminary_selection.xlsx",
    "comparison_01": "02_comparison_01_metric_utility.xlsx",
    "comparison_02": "03_comparison_02_meta_learning.xlsx",
    "comparison_03": "04_comparison_03_raw_softmax_vbs.xlsx",
    "comparison_04": "05_comparison_04_operational_benchmarks.xlsx",
    "comparison_05": "06_comparison_05_vbs_headroom.xlsx",
    "metric_selection": "07_metric_selection.xlsx",
    "gap_decomposition": "08_gap_decomposition.xlsx",
    "runtime_pareto": "09_runtime_and_pareto.xlsx",
    "statistical_tests": "10_statistical_tests.xlsx",
    "full_analysis": "11_full_analysis.xlsx",
}
